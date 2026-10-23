"""Import explicit source values only; do not infer deadlines or completion."""
import json, sys, datetime
from pathlib import Path
import openpyxl

def convert(path):
    book = openpyxl.load_workbook(path, data_only=True)
    sites = []
    def clean(value):
        if value is None: return ''
        if isinstance(value, float) and value.is_integer(): return str(int(value))
        return ' '.join(str(value).split())
    for sheet in book:
        if sheet.cell(1, 1).value != 'SORUMLU' or not sheet.cell(2, 3).value: continue
        tasks = []
        for row in range(3, sheet.max_row + 1):
            text = clean(sheet.cell(row, 2).value)
            if text:
                tasks.append(dict(id=f'{len(sites)+1}-{row}', text=text, owner=clean(sheet.cell(row,1).value) or 'Atanmamış', row=row))
        sites.append(dict(id=str(len(sites)+1), name=clean(sheet.cell(2,3).value), ada=clean(sheet.cell(2,4).value), parcel=clean(sheet.cell(2,5).value), contact=clean(sheet.cell(2,6).value), area=clean(sheet.cell(2,7).value), units=clean(sheet.cell(2,8).value), period=clean(sheet.cell(2,9).value), sheet=sheet.title, tasks=tasks))
    # Match overview order, but retain details from individual sheets.
    overview=book['Tüm İşler']
    order={}
    for col in range(3,overview.max_column+1,2):
        order[(clean(overview.cell(2,col).value),clean(overview.cell(3,col).value))]=col
    for site in sites:
        col = order.get((site['ada'],site['parcel']))
        if col is not None:
            site['inspection'] = clean(overview.cell(5,col).value)
    sites.sort(key=lambda s:order.get((s['ada'],s['parcel']),999))
    return dict(source=Path(path).name, importedAt=datetime.datetime.now(datetime.timezone.utc).isoformat(), sites=sites)
if __name__=='__main__':
    result=convert(sys.argv[1])
    out=Path(__file__).resolve().parents[1]/'web/data.json'
    out.write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(f"{len(result['sites'])} şantiye, {sum(len(s['tasks']) for s in result['sites'])} kayıt aktarıldı.")
